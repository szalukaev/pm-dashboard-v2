"""Generate invoice_template.docx for PM Dashboard V3."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shading_elm = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color_hex,
    })
    shading.append(shading_elm)

def set_cell_border(cell, **kwargs):
    """Set cell borders."""
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.makeelement(qn('w:tcBorders'), {})
    for edge, val in kwargs.items():
        element = tcBorders.makeelement(qn(f'w:{edge}'), {
            qn('w:val'): val.get('val', 'single'),
            qn('w:sz'): val.get('sz', '4'),
            qn('w:color'): val.get('color', '000000'),
            qn('w:space'): '0',
        })
        tcBorders.append(element)
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text, font_name='Calibri', font_size=11,
                         bold=False, alignment=None, color=None, space_after=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if alignment:
        p.alignment = alignment
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p

def create_invoice_template():
    doc = Document()

    # Page setup — A4
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    # Default font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # ═══════════════════════════════════════════
    # TITLE
    # ═══════════════════════════════════════════
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('СЧЁТ НА ОПЛАТУ')
    run.font.name = 'Calibri Light'
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p.paragraph_format.space_after = Pt(4)

    # Invoice number and date line
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run('№ ')
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.bold = True
    run = p.add_run('{{invoice_num}}')
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.underline = True

    run = p.add_run('        ')
    run = p.add_run('Дата: ')
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.bold = True
    run = p.add_run('{{date}}')
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.underline = True

    run = p.add_run('        ')
    run = p.add_run('{{vat_label}}')
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor(0x5E, 0x6A, 0xD2)
    p.paragraph_format.space_after = Pt(12)

    # ═══════════════════════════════════════════
    # RECIPIENT TABLE
    # ═══════════════════════════════════════════
    add_styled_paragraph(doc, 'Получатель:', font_size=10, bold=True,
                         color=RGBColor(0x59, 0x59, 0x59), space_after=4)

    table0 = doc.add_table(rows=4, cols=2)
    table0.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Set column widths
    for row in table0.rows:
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(12.0)

    labels = ['Название', 'Адрес', 'Телефон', 'Контактное лицо']
    placeholders = ['{{company}}', '{{address}}', '{{phone}}', '{{contact}}']

    for i, (label, placeholder) in enumerate(zip(labels, placeholders)):
        cell_label = table0.rows[i].cells[0]
        cell_value = table0.rows[i].cells[1]

        # Label cell
        p = cell_label.paragraphs[0]
        run = p.add_run(label)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        set_cell_shading(cell_label, 'F5F6F6')

        # Value cell
        p = cell_value.paragraphs[0]
        run = p.add_run(placeholder)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.underline = True

    # Table borders
    for row in table0.rows:
        for cell in row.cells:
            set_cell_border(cell,
                top={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
                bottom={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
                left={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
                right={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
            )

    doc.add_paragraph()  # spacer

    # ═══════════════════════════════════════════
    # SERVICES TABLE
    # ═══════════════════════════════════════════
    add_styled_paragraph(doc, 'Наименование работ и услуг:', font_size=10, bold=True,
                         color=RGBColor(0x59, 0x59, 0x59), space_after=4)

    table1 = doc.add_table(rows=6, cols=5)
    table1.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Column widths
    col_widths = [Cm(1.2), Cm(7.5), Cm(2.0), Cm(3.0), Cm(3.0)]
    for row in table1.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # Header row
    headers = ['№', 'Наименование', 'Кол-во', 'Цена', 'Сумма']
    for i, h in enumerate(headers):
        cell = table1.rows[0].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, '1F3A5F')

    # Template data row
    data_row = ['1', '{{item_name}}', '{{item_qty}}', '{{item_price}}', '{{item_amount}}']
    for i, val in enumerate(data_row):
        cell = table1.rows[1].cells[i]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i != 1 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(val)
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.underline = True

    # Empty rows 2-3 (for additional items)
    for r in range(2, 4):
        for c in range(5):
            cell = table1.rows[r].cells[c]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('')
            run.font.size = Pt(10)

    # Subtotal row
    subtotal_row_idx = 4
    cell_sub_label = table1.rows[subtotal_row_idx].cells[3]
    p = cell_sub_label.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('Итого:')
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.bold = True
    set_cell_shading(cell_sub_label, 'F5F6F6')

    cell_sub_value = table1.rows[subtotal_row_idx].cells[4]
    p = cell_sub_value.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('{{subtotal}}')
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.bold = True
    run.underline = True
    set_cell_shading(cell_sub_value, 'F5F6F6')

    # VAT row
    vat_row_idx = 5
    cell_vat_label = table1.rows[vat_row_idx].cells[3]
    p = cell_vat_label.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('{{vat_label}}:')
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    set_cell_shading(cell_vat_label, 'F5F6F6')

    cell_vat_value = table1.rows[vat_row_idx].cells[4]
    p = cell_vat_value.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('{{vat_amount}}')
    run.font.name = 'Calibri'
    run.font.size = Pt(10)
    run.underline = True
    set_cell_shading(cell_vat_value, 'F5F6F6')

    # Table borders
    for row in table1.rows:
        for cell in row.cells:
            set_cell_border(cell,
                top={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
                bottom={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
                left={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
                right={'val': 'single', 'sz': '4', 'color': 'E0E2E6'},
            )

    doc.add_paragraph()  # spacer

    # ═══════════════════════════════════════════
    # TOTAL
    # ═══════════════════════════════════════════
    total_table = doc.add_table(rows=1, cols=2)
    total_table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    total_table.rows[0].cells[0].width = Cm(5.0)
    total_table.rows[0].cells[1].width = Cm(4.0)

    cell_label = total_table.rows[0].cells[0]
    p = cell_label.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('ВСЕГО К ОПЛАТЕ:')
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    set_cell_shading(cell_label, 'E8EDF3')

    cell_value = total_table.rows[0].cells[1]
    p = cell_value.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run('{{total}}')
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    run.underline = True
    set_cell_shading(cell_value, 'E8EDF3')

    for cell in total_table.rows[0].cells:
        set_cell_border(cell,
            top={'val': 'single', 'sz': '6', 'color': '1F3A5F'},
            bottom={'val': 'single', 'sz': '6', 'color': '1F3A5F'},
            left={'val': 'single', 'sz': '6', 'color': '1F3A5F'},
            right={'val': 'single', 'sz': '6', 'color': '1F3A5F'},
        )

    doc.add_paragraph()  # spacer

    # ═══════════════════════════════════════════
    # FOOTER NOTE
    # ═══════════════════════════════════════════
    p = doc.add_paragraph()
    run = p.add_run('Счёт сформирован автоматически системой PM Dashboard.')
    run.font.name = 'Calibri'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    run.italic = True
    p.paragraph_format.space_before = Pt(20)

    # Save
    out_path = '/Users/sergei/YDisk/Projects/pm-dashboard-V3/backend/invoice/invoice_template.docx'
    doc.save(out_path)
    print(f'Template saved to: {out_path}')

if __name__ == '__main__':
    create_invoice_template()
